#!/usr/bin/env python3
"""~/Downloads/nya-bilder/<id>.<valfritt format> -> img/<id>.jpg (400px, vit bakgrund) + Excel: image=<id>.jpg, active=ja.
Kör från repo-roten: python3 tools/nya-bilder.py   (--kör = gör det på riktigt)"""
import os, sys, shutil, subprocess, tempfile
from openpyxl import load_workbook
from PIL import Image
KOR = '--kör' in sys.argv
SRC = os.path.expanduser('~/Downloads/nya-bilder'); XLSX = 'data/snuskampen-master.xlsx'
wb = load_workbook(XLSX); ws = wb['Produkter']
cols = [c.value for c in ws[1]]; ci, ca, cm = (cols.index(k) + 1 for k in ('id', 'active', 'image'))
row = {ws.cell(r, ci).value: r for r in range(2, ws.max_row + 1) if ws.cell(r, ci).value}

def open_any(p):
    try: return Image.open(p)
    except Exception:
        tmp = tempfile.mktemp(suffix='.png')   # AVIF/HEIC m.m. via macOS sips
        subprocess.run(['sips', '-s', 'format', 'png', p, '--out', tmp], check=True, capture_output=True)
        return Image.open(tmp)

klara, fel = [], []
for f in sorted(os.listdir(SRC)):
    pid, ext = os.path.splitext(f)
    if f.startswith('.') or ext.lower() not in ('.jpg', '.jpeg', '.png', '.webp', '.avif', '.heic', '.gif'): continue
    if pid not in row: fel.append(f'okänt id: {f}'); continue
    try: im = open_any(os.path.join(SRC, f))
    except Exception as e: fel.append(f'kunde inte läsa {f}: {e}'); continue
    im = im.convert('RGBA'); bg = Image.new('RGBA', im.size, (255, 255, 255, 255)); bg.alpha_composite(im)
    im = bg.convert('RGB'); im.thumbnail((400, 400), Image.LANCZOS)
    sq = Image.new('RGB', (400, 400), 'white'); sq.paste(im, ((400 - im.width) // 2, (400 - im.height) // 2))
    r = row[pid]; was = (ws.cell(r, cm).value, ws.cell(r, ca).value)
    klara.append(f'{pid}  (bild: {was[0] or "–"} → {pid}.jpg, active: {was[1]} → ja)')
    if KOR:
        sq.save(f'img/{pid}.jpg', quality=90)
        ws.cell(r, cm).value = f'{pid}.jpg'; ws.cell(r, ca).value = 'ja'
for k in klara: print('  ok  ', k)
for e in fel: print('  FEL ', e)
print(f'\n{len(klara)} klara, {len(fel)} fel')
if KOR and klara: wb.save(XLSX); print('Excel sparad, bilder skrivna till img/. Kör nu build.py.')
elif not KOR: print('Torrkörning – lägg till --kör')
