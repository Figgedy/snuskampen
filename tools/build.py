#!/usr/bin/env python3
"""data/snuskampen-master.xlsx -> products.js + images.js, med kontroller.
Kör från repo-roten: python3 tools/build.py   (--dry = kontrollera utan att skriva)"""
import json, re, os, sys
from collections import Counter, OrderedDict
from openpyxl import load_workbook

SRC = 'data/snuskampen-master.xlsx'
DRY = '--dry' in sys.argv
BASE = ['id', 'brand', 'maker', 'name', 'format', 'style', 'flavor', 'mg', 'dots', 'type']
TRUE, FALSE = {'ja', 'yes', 'true', 'sant', '1', 'x'}, {'nej', 'no', 'false', 'falskt', '0', ''}

ws = load_workbook(SRC, data_only=True)['Produkter']
rows = list(ws.iter_rows(values_only=True))
cols = [str(c).strip() if c is not None else '' for c in rows[0]]
errors, warns, prods, imgs = [], [], [], OrderedDict()

def clean(v):
    if v is None: return ''
    if isinstance(v, float) and v.is_integer(): return int(v)
    return v.strip() if isinstance(v, str) else v

for n, r in enumerate(rows[1:], start=2):
    d = {c: clean(v) for c, v in zip(cols, r) if c}
    if not any(d.get(c) not in ('', None) for c in ('id', 'brand', 'name')): continue
    pid = str(d.get('id', ''))
    tag = f'rad {n} ({pid or "utan id"})'
    if not pid: errors.append(f'{tag}: id saknas'); continue
    if not re.fullmatch(r'[a-z0-9][a-z0-9-]*', pid): errors.append(f'{tag}: id får bara ha gemener, siffror och bindestreck')
    for req in ('brand', 'name'):
        if not str(d.get(req, '')).strip(): errors.append(f'{tag}: {req} saknas')
    act = str(d.get('active', 'ja')).strip().lower()
    if act not in TRUE | FALSE: errors.append(f'{tag}: active måste vara ja/nej (är "{d.get("active")}")')
    active = act in TRUE
    obj = OrderedDict()
    for c in BASE:
        if c in d and d[c] != '': obj[c] = d[c]
    for c in cols:
        if c and c not in BASE and c not in ('active', 'image', 'kommentar') and d.get(c, '') != '':
            v = d[c]
            if isinstance(v, str) and v[:1] in '[{':
                try: v = json.loads(v)
                except ValueError: pass
            obj[c] = v
    if 'dots' in obj:
        try: obj['dots'] = int(obj['dots'])
        except (TypeError, ValueError): errors.append(f'{tag}: dots måste vara ett heltal (är "{obj["dots"]}")')
    if not active: obj['active'] = False
    img = str(d.get('image', '') or '').strip()
    if img:
        if not os.path.exists(os.path.join('img', img)): errors.append(f'{tag}: bildfilen img/{img} finns inte')
        imgs[pid] = img
    elif active: warns.append(f'{tag}: aktiv men saknar bild')
    if str(obj.get('brand', '')).lower() and str(obj.get('name', '')).lower().startswith(str(obj.get('brand', '')).lower() + ' '):
        warns.append(f'{tag}: namnet börjar med märket ("{obj["brand"]} {obj["name"]}")')
    prods.append(obj)

dup = [i for i, c in Counter(p['id'] for p in prods).items() if c > 1]
for i in dup: errors.append(f'dubblett-id: {i}')
seen = Counter((str(p.get('brand', '')).lower(), str(p.get('name', '')).lower()) for p in prods)
for (b, nm), c in seen.items():
    if c > 1: warns.append(f'samma märke+namn {c} gånger: {b} {nm}')
for field in ('style', 'format', 'flavor', 'type', 'brand'):
    cnt = Counter(str(p.get(field, '')) for p in prods if p.get(field, '') != '')
    low = {k: v for k, v in cnt.items() if v == 1}
    if low and len(cnt) > 3:
        for k in low:
            near = [o for o in cnt if o != k and o.lower()[:3] == k.lower()[:3]]
            if near: warns.append(f'{field} "{k}" förekommer bara en gång – stavfel för {near}?')

print(f'{len(prods)} produkter · {sum(1 for p in prods if p.get("active") is not False)} aktiva · {len(imgs)} med bild')
for w in warns: print('  varning:', w)
if errors:
    for e in errors: print('  FEL:', e)
    sys.exit(f'{len(errors)} fel – inget skrivet. Rätta i Excel och kör igen.')
if DRY: sys.exit('Torrkörning: inga fel. Inget skrivet.')

out = ['// SNUSKAMPEN – produktdata. GENERERAD från data/snuskampen-master.xlsx av tools/build.py.',
       '// Redigera inte här – ändra i Excel och kör python3 tools/build.py.',
       'window.PRODUCTS = [']
last = None
for p in prods:
    if p.get('brand') != last: out.append(f'  // {p.get("brand")}'); last = p.get('brand')
    out.append('  ' + json.dumps(p, ensure_ascii=False) + ',')
out.append('];')
open('products.js', 'w', encoding='utf-8').write('\n'.join(out) + '\n')
im = ['// id -> filnamn i /img/. GENERERAD av tools/build.py.', 'window.IMAGES = {']
im.append(',\n'.join(f'{json.dumps(k, ensure_ascii=False)}: {json.dumps(v, ensure_ascii=False)}' for k, v in imgs.items()))
im.append('};')
open('images.js', 'w', encoding='utf-8').write('\n'.join(im) + '\n')
print('products.js och images.js skrivna.')
