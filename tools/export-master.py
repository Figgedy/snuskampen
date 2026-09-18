#!/usr/bin/env python3
"""Engångsexport: products.js + images.js -> data/snuskampen-master.xlsx
Kör från repo-roten: python3 tools/export-master.py"""
import json, re, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUT = 'data/snuskampen-master.xlsx'
if os.path.exists(OUT) and '--force' not in sys.argv:
    sys.exit(f'{OUT} finns redan. Kör med --force om du verkligen vill skriva över den.')

def load_products(path='products.js'):
    s = open(path, encoding='utf-8').read()
    body = s[s.index('['):s.rindex(']') + 1]
    body = '\n'.join(l for l in body.split('\n') if not l.strip().startswith('//'))
    body = re.sub(r',\s*\]', ']', body)
    return json.loads(body)

def load_images(path='images.js'):
    s = open(path, encoding='utf-8').read()
    return dict(re.findall(r'^\s*"([^"]+)"\s*:\s*"([^"]+)"', s, re.M))

prods, imgs = load_products(), load_images()
BASE = ['id', 'active', 'brand', 'name', 'maker', 'format', 'style', 'flavor', 'mg', 'dots', 'type']
extra = []
for p in prods:
    for k in p:
        if k not in BASE and k not in extra: extra.append(k)
cols = BASE + extra + ['image', 'kommentar']

wb = Workbook(); ws = wb.active; ws.title = 'Produkter'
ws.append(cols)
for p in prods:
    row = []
    for c in cols:
        if c == 'active': row.append('nej' if p.get('active') is False else 'ja')
        elif c == 'image': row.append(imgs.get(p['id'], ''))
        elif c == 'kommentar': row.append('')
        else:
            v = p.get(c, '')
            row.append(json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v)
    ws.append(row)

hdr = Font(bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='111820')
for i, c in enumerate(cols, 1):
    cell = ws.cell(1, i); cell.font = hdr; cell.fill = fill; cell.alignment = Alignment(vertical='center')
    w = max(len(str(c)), *(len(str(ws.cell(r, i).value or '')) for r in range(2, min(ws.max_row, 400) + 1)))
    ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 8), 48)
ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions
lock = PatternFill('solid', fgColor='EFEFEF')
for r in range(2, ws.max_row + 1): ws.cell(r, 1).fill = lock

info = wb.create_sheet('Info')
for line in [
    ['Kolumn', 'Förklaring'],
    ['id', 'ÄNDRA ALDRIG på befintliga produkter – röster och delade länkar hänger på id. Nya produkter: gemener, bindestreck, t.ex. zyn-cool-mint-slim-6.'],
    ['active', 'ja = syns i dueller och topplistor. nej = dold (vs-länkar fungerar ändå).'],
    ['brand / name', 'Visas på sajten som "brand name". Skriv inte märket i name.'],
    ['format', 'Storlek: Mini, Slim, Large, Normal osv.'],
    ['style', 'Typ: White, Original, White Dry osv. Styr filter och månadskategorier.'],
    ['flavor', 'Smakfamilj på svenska. Styr smakfiltret.'],
    ['mg', 'Fritext som visas, t.ex. "6 mg/prilla · 10 mg/g".'],
    ['dots', 'Styrka 1–5 (prickar).'],
    ['type', 'tobak eller nikotin.'],
    ['image', 'Filnamn i img/. Tomt = ingen bild.'],
    ['kommentar', 'Fritt för dig. Följer inte med till sajten.'],
    ['', ''],
    ['Arbetsflöde', 'Redigera -> spara -> python3 tools/build.py -> git add -A && git commit -m "..." && git push'],
]: info.append(line)
info.column_dimensions['A'].width = 16; info.column_dimensions['B'].width = 110
info['A1'].font = Font(bold=True); info['B1'].font = Font(bold=True)

os.makedirs('data', exist_ok=True); wb.save(OUT)
print(f'{OUT}: {len(prods)} produkter, {sum(1 for p in prods if p["id"] in imgs)} med bild, kolumner: {", ".join(cols)}')
